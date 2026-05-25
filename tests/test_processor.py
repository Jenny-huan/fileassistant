from __future__ import annotations

from io import BytesIO
from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from PIL import Image

from app.config import MAX_UPLOAD_BYTES
from app.docx_inspector import get_body_width_emu
from app.processor import ProcessingError, inspect_upload, process_existing_job, process_upload


class ProcessorTests(unittest.TestCase):
    def test_rejects_non_docx(self) -> None:
        with self.assertRaises(ProcessingError):
            process_upload("bad.txt", b"not a docx")

    def test_rejects_empty_docx(self) -> None:
        with self.assertRaises(ProcessingError):
            process_upload("empty.docx", b"")

    def test_rejects_oversized_docx(self) -> None:
        with self.assertRaises(ProcessingError):
            process_upload("big.docx", b"x" * (MAX_UPLOAD_BYTES + 1))

    def test_formats_and_reports_docx_images(self) -> None:
        source = _build_sample_docx()
        report = process_upload("sample.docx", source.read_bytes())

        self.assertIn("job_id", report)
        self.assertTrue(Path(report["backup_file"]).exists())
        self.assertTrue(Path(report["output_file"]).exists())
        self.assertEqual(report["stats"]["inline_images"], 2)
        self.assertEqual(report["stats"]["media_images"], 2)
        self.assertEqual(report["stats"]["resized_images"], 2)
        self.assertGreaterEqual(report["stats"]["centered_paragraphs"], 2)

        result = Document(report["output_file"])
        self.assertEqual(len(result.inline_shapes), 2)
        max_width = int(get_body_width_emu(result) * 0.8)
        self.assertEqual(int(result.inline_shapes[0].width), max_width)
        self.assertEqual(int(result.inline_shapes[1].width), max_width)

        image_paragraphs = [p for p in result.paragraphs if p._p.xpath(".//w:drawing")]
        self.assertEqual(len(image_paragraphs), 2)
        for paragraph in image_paragraphs:
            self.assertEqual(paragraph.alignment, WD_ALIGN_PARAGRAPH.CENTER)

    def test_inspect_then_processes_with_custom_ratio(self) -> None:
        source = _build_sample_docx()
        inspected = inspect_upload("sample.docx", source.read_bytes())
        self.assertIn("job_id", inspected)
        self.assertEqual(inspected["document"]["inline_images"], 2)
        self.assertEqual(inspected["settings"]["default_width_ratio"], 0.8)

        report = process_existing_job(inspected["job_id"], max_width_ratio=0.6, alignment="right")
        self.assertEqual(report["job_id"], inspected["job_id"])
        self.assertEqual(report["settings"]["max_width_ratio"], 0.6)
        self.assertEqual(report["settings"]["alignment"], "right")
        self.assertEqual(report["stats"]["resized_images"], 2)

        result = Document(report["output_file"])
        expected_width = int(get_body_width_emu(result) * 0.6)
        self.assertEqual(int(result.inline_shapes[0].width), expected_width)
        self.assertEqual(int(result.inline_shapes[1].width), expected_width)
        image_paragraphs = [p for p in result.paragraphs if p._p.xpath(".//w:drawing")]
        for paragraph in image_paragraphs:
            self.assertEqual(paragraph.alignment, WD_ALIGN_PARAGRAPH.RIGHT)

    def test_rejects_invalid_alignment(self) -> None:
        source = _build_sample_docx()
        with self.assertRaises(ProcessingError):
            process_upload("sample.docx", source.read_bytes(), alignment="justify")


def _build_sample_docx() -> Path:
    tmp = Path(tempfile.mkdtemp())
    wide = tmp / "wide.jpg"
    small = tmp / "small.png"
    docx_path = tmp / "sample.docx"

    Image.new("RGB", (2400, 1200), "#3478a6").save(wide, "JPEG", quality=95)
    Image.new("RGBA", (400, 240), (255, 0, 0, 96)).save(small, "PNG")

    document = Document()
    document.add_paragraph("Before image")
    document.add_picture(str(wide), width=Inches(7))
    document.add_paragraph("Between images")
    document.add_picture(str(small), width=Inches(1))
    document.save(docx_path)
    return docx_path


if __name__ == "__main__":
    unittest.main()
